from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).parent
DOCS = ROOT / "docs"
DOCS.mkdir(exist_ok=True)
SCREENSHOT = DOCS / "dashboard-screenshot.png"
OUTPUT = DOCS / "csv-user-manager-presentation.pptx"

NAVY = (33, 37, 41)
BLUE = (78, 115, 223)
HOST_BLUE = (46, 89, 217)
LIGHT = (248, 249, 250)
GRAY = (108, 117, 125)

def font(size, bold=False):
    name = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    return ImageFont.truetype(name, size)

def dashboard_image():
    im = Image.new("RGB", (1600, 980), LIGHT)
    d = ImageDraw.Draw(im)
    d.rectangle((0, 0, 1600, 62), fill=NAVY)
    d.text((30, 17), "System Infrastructure Database Dashboard", fill="white", font=font(25, True))
    d.rounded_rectangle((1450, 14, 1565, 49), 5, fill=(220, 53, 69))
    d.text((1470, 22), "Logout", fill="white", font=font(16, True))
    panels = [(28, 88, 782, 930, BLUE, "User Management"), (818, 88, 1572, 930, HOST_BLUE, "Host Access Management")]
    for x1, y1, x2, y2, color, title in panels:
        d.rounded_rectangle((x1, y1, x2, y2), 10, fill="white", outline=color, width=4)
        d.text((x1 + 18, y1 + 16), title, fill=color, font=font(25, True))
    d.rounded_rectangle((52, 135, 758, 390), 8, fill="white", outline=BLUE, width=3)
    d.text((72, 150), "Add User", fill=BLUE, font=font(21, True))
    fields = ["Username", "Email", "UID", "GID", "Home Directory", "Public Key", "Authorized Host"]
    y = 190
    for label in fields:
        d.text((75, y), label, fill=(73, 80, 87), font=font(15, True))
        d.rounded_rectangle((225, y - 3, 730, y + 25), 4, fill=(250, 250, 250), outline=(206, 212, 218))
        y += 28
    d.rounded_rectangle((52, 416, 758, 880), 8, fill="white", outline=BLUE, width=3)
    d.text((72, 432), "Users List (users.csv)", fill=BLUE, font=font(21, True))
    d.rounded_rectangle((72, 470, 500, 503), 4, fill=(250, 250, 250), outline=(206, 212, 218))
    d.text((82, 477), "Search Users", fill=GRAY, font=font(14))
    headers = ["User", "UID/GID", "Email", "Auth Host", "Actions"]
    xs = [72, 215, 325, 505, 650]
    for x, h in zip(xs, headers):
        d.text((x, 530), h, fill=BLUE, font=font(14, True))
    rows = [("tsandholm", "3000/3000", "tom@...", "*"), ("kat", "3001/3001", "kat@...", "*"), ("mary", "3002/3002", "mary@...", "tom2")]
    y = 568
    for row in rows:
        for x, value in zip(xs, row):
            d.text((x, y), value, fill=(55, 55, 55), font=font(14))
        d.rounded_rectangle((650, y - 3, 730, y + 22), 4, fill=(23, 162, 184))
        d.text((666, y + 2), "View", fill="white", font=font(12, True))
        y += 42
    d.rounded_rectangle((842, 135, 1548, 390), 8, fill="white", outline=HOST_BLUE, width=3)
    d.text((862, 150), "Add New Host Group", fill=HOST_BLUE, font=font(21, True))
    for i, label in enumerate(["Machine-Group", "Group ID", "Member List"]):
        yy = 200 + i * 48
        d.text((865, yy), label, fill=(73, 80, 87), font=font(15, True))
        d.rounded_rectangle((1025, yy - 3, 1520, yy + 27), 4, fill=(250, 250, 250), outline=(206, 212, 218))
    d.rounded_rectangle((842, 416, 1548, 880), 8, fill="white", outline=HOST_BLUE, width=3)
    d.text((862, 432), "Hosts List (hosts.csv)", fill=HOST_BLUE, font=font(21, True))
    for x, h in zip([862, 1110, 1225, 1435], ["Machine-Group", "Group ID", "Members", "Actions"]):
        d.text((x, 492), h, fill=HOST_BLUE, font=font(14, True))
    for i, row in enumerate([("tom1-tsand-org", "5000", "ansible,sudo,tom"), ("tom2-tsand-org", "5001", "ansible,sudo,mary"), ("tom3-tsand-org", "5002", "ansible,sudo,kat")]):
        yy = 535 + i * 48
        for x, value in zip([862, 1110, 1225], row):
            d.text((x, yy), value, fill=(55, 55, 55), font=font(14))
    d.rounded_rectangle((842, 730, 1548, 875), 8, fill=(248, 249, 252), outline=HOST_BLUE, width=2)
    d.text((862, 745), "Next Step: Run Ansible Playbooks", fill=HOST_BLUE, font=font(18, True))
    d.text((862, 780), "As user ansible, run in order:", fill=(55, 55, 55), font=font(14))
    for i, item in enumerate(["update-group-block.yml", "update-user-block.yml", "setup-local-user-homes.yml"]):
        d.text((885, 805 + i * 20), f"{i+1}. {item}", fill=(55, 55, 55), font=font(13))
    im.save(SCREENSHOT)

def add_textbox(slide, text, x, y, w, h, size=20, color=(40, 40, 40), bold=False, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Aptos"
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = RGBColor(*color)
    if align:
        p.alignment = align
    return box

def title(slide, text, subtitle=None):
    add_textbox(slide, text, 0.6, 0.35, 12.1, 0.55, 28, NAVY, True)
    if subtitle:
        add_textbox(slide, subtitle, 0.62, 0.95, 12, 0.35, 13, GRAY)

def bullets(slide, items, x=0.9, y=1.5, w=11.3, size=21, gap=0.48):
    for i, item in enumerate(items):
        add_textbox(slide, "\u2022 " + item, x, y + i * gap, w, 0.38, size, (55, 55, 55))

def add_card(slide, x, y, w, h, heading, body, color=BLUE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape.line.color.rgb = RGBColor(*color); shape.line.width = Pt(1.5)
    add_textbox(slide, heading, x + 0.18, y + 0.15, w - 0.35, 0.35, 18, color, True)
    add_textbox(slide, body, x + 0.18, y + 0.62, w - 0.35, h - 0.75, 14, (65, 65, 65))

dashboard_image()
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

slide = prs.slides.add_slide(blank)
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor(*NAVY)
add_textbox(slide, "CSV User Manager", 0.75, 1.35, 11.8, 0.8, 38, (255,255,255), True)
add_textbox(slide, "A focused dashboard for CSV-backed identity and host access management", 0.8, 2.25, 11.5, 0.5, 22, (210,220,235))
add_textbox(slide, "Project overview and operational workflow", 0.8, 5.85, 11.5, 0.4, 16, (170,185,205))

slide = prs.slides.add_slide(blank); title(slide, "Project goals", "Make small-scale identity and access administration predictable, visible, and repeatable.")
bullets(slide, ["Manage users and machine-groups from one authenticated PHP dashboard.", "Keep CSV files human-readable and easy to review or version.", "Generate deployment-ready passwd and group blocks from source data.", "Apply changes deliberately through explicit, manual Ansible playbooks.", "Protect unrelated host membership data when one user assignment changes."])

slide = prs.slides.add_slide(blank); title(slide, "Dashboard at a glance", "The UI separates user identity management from host access management.")
slide.shapes.add_picture(str(SCREENSHOT), Inches(0.55), Inches(1.35), width=Inches(12.25))

slide = prs.slides.add_slide(blank); title(slide, "Core features")
add_card(slide, 0.7, 1.35, 3.8, 1.55, "User management", "Add, edit, delete, search, and inspect complete user records. Assign a specific Auth Host or all hosts.", BLUE)
add_card(slide, 4.8, 1.35, 3.8, 1.55, "Host management", "Maintain machine-groups, IDs, and member lists while preserving existing members during targeted updates.", HOST_BLUE)
add_card(slide, 8.9, 1.35, 3.8, 1.55, "Raw file controls", "View or edit source CSVs and generated blocks in authenticated, readable pages.", BLUE)
add_card(slide, 0.7, 3.35, 3.8, 1.55, "Safe publishing", "Publish hosts-block.txt separately from users-block.txt and groups-block.txt; publishing never runs Ansible.", HOST_BLUE)
add_card(slide, 4.8, 3.35, 3.8, 1.55, "Operational safeguards", "Password hashing, locked raw-file writes, escaped output, atomic user-block generation, and backups.", BLUE)
add_card(slide, 8.9, 3.35, 3.8, 1.55, "Usable layout", "Side-by-side management areas, searchable lists, internal list scrolling, and page-level scrolling.", HOST_BLUE)

slide = prs.slides.add_slide(blank); title(slide, "Project structure", "Small, explicit files keep the application easy to deploy and understand.")
add_card(slide, 0.7, 1.35, 3.8, 2.0, "Web application", "`index.php` handles authentication, CSV operations, synchronization, publishing, and raw views. `view.php` renders the dashboard.", BLUE)
add_card(slide, 4.8, 1.35, 3.8, 2.0, "Source data", "`users.csv` stores identity records. `hosts.csv` stores machine-groups and member lists.", HOST_BLUE)
add_card(slide, 8.9, 1.35, 3.8, 2.0, "Generated artifacts", "`users-block.txt`, `groups-block.txt`, and `hosts-block.txt` translate dashboard data into system-file formats.", BLUE)
add_card(slide, 2.75, 4.05, 3.8, 1.65, "Remote automation", "`update-group-block.yml` and `update-user-block.yml` update remote `/etc/group` and `/etc/passwd` on `virt` hosts.", HOST_BLUE)
add_card(slide, 6.85, 4.05, 3.8, 1.65, "Controller automation", "`setup-local-user-homes.yml` runs only on localhost to create homes and install SSH public keys.", BLUE)

slide = prs.slides.add_slide(blank); title(slide, "Data-to-system flow", "Source files remain the authority; generated blocks are reviewed before manual application.")
steps = [("1", "Edit CSVs", "Dashboard or raw editor"), ("2", "Publish blocks", "Generate passwd/group files"), ("3", "Deploy files", "`make push` to /var/www/html"), ("4", "Apply remotely", "Ansible with check/diff"), ("5", "Prepare controller", "Homes and authorized keys")]
for i, (n, head, body) in enumerate(steps):
    x = 0.55 + i * 2.55
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.2), Inches(2.15), Inches(1.7))
    shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor(248,249,252)
    shape.line.color.rgb = RGBColor(*(BLUE if i % 2 == 0 else HOST_BLUE))
    add_textbox(slide, n, x+0.15, 2.38, 0.35, 0.35, 22, BLUE if i % 2 == 0 else HOST_BLUE, True)
    add_textbox(slide, head, x+0.55, 2.38, 1.45, 0.35, 16, NAVY, True)
    add_textbox(slide, body, x+0.15, 2.95, 1.85, 0.65, 13, (65,65,65))
    if i < 4: add_textbox(slide, ">", x+2.23, 2.75, 0.3, 0.4, 22, GRAY, True, PP_ALIGN.CENTER)

slide = prs.slides.add_slide(blank); title(slide, "Ansible playbooks and order", "Run as user `ansible` from the controller; preview with `--check --diff` first.")
add_card(slide, 0.7, 1.35, 3.75, 3.8, "1. update-group-block.yml", "Target: remote `virt` hosts\n\nReads `hosts-block.txt` and updates the managed section of `/etc/group` with backups and privilege escalation.", HOST_BLUE)
add_card(slide, 4.8, 1.35, 3.75, 3.8, "2. update-user-block.yml", "Target: remote `virt` hosts\n\nCreates primary groups, updates the managed `/etc/passwd` section, and runs `pwconv`.", BLUE)
add_card(slide, 8.9, 1.35, 3.75, 3.8, "3. setup-local-user-homes.yml", "Target: `localhost` only\n\nCreates `/share/home` directories and `.ssh/authorized_keys` with correct ownership and permissions.", HOST_BLUE)
add_textbox(slide, "Recommended commands: ansible-playbook -i inventory update-group-block.yml --check --diff  →  apply  →  update-user-block.yml  →  setup-local-user-homes.yml", 0.85, 5.75, 11.7, 0.65, 15, NAVY, True)

slide = prs.slides.add_slide(blank); title(slide, "Advantages and operating principles")
bullets(slide, ["Simple deployment model: PHP, CSV, generated blocks, and playbooks can live in `/var/www/html`.", "Human-readable inputs support review, backup, and straightforward recovery.", "Manual Ansible execution keeps production changes intentional and observable.", "Targeted host synchronization preserves unrelated members instead of rebuilding lists.", "Generated files use standard `/etc/passwd` and `/etc/group` formats.", "Authentication and one-way password hashing protect dashboard access."])

slide = prs.slides.add_slide(blank); title(slide, "Quick start")
add_textbox(slide, "1. Start the dashboard", 0.9, 1.45, 3.5, 0.35, 21, BLUE, True)
add_textbox(slide, "php -S 127.0.0.1:8000", 0.9, 1.9, 5.2, 0.45, 19, NAVY, True)
add_textbox(slide, "2. Edit and publish", 0.9, 2.75, 3.5, 0.35, 21, HOST_BLUE, True)
add_textbox(slide, "Save users/hosts, then publish hosts-block.txt and users-block.txt.", 0.9, 3.2, 5.8, 0.45, 17, (55,55,55))
add_textbox(slide, "3. Apply in order", 0.9, 4.05, 3.5, 0.35, 21, BLUE, True)
add_textbox(slide, "update-group-block.yml\nupdate-user-block.yml\nsetup-local-user-homes.yml", 0.9, 4.5, 5.8, 1.2, 18, NAVY, True)
add_textbox(slide, "The dashboard provides the data and review surface; Ansible performs the controlled system changes.", 7.0, 2.0, 5.2, 1.4, 24, NAVY, True)

prs.save(OUTPUT)
print(OUTPUT)
